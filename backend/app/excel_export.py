"""Builds a formatted Excel workbook from generated test cases."""
from __future__ import annotations

from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import TestCase

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PRIORITY_FILLS = {
    "P1": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    "P2": PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid"),
    "P3": PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid"),
}

COLUMNS = [
    ("tc_id", "TC ID", 12),
    ("module", "Module", 24),
    ("case_type", "Type", 12),
    ("priority", "Priority", 10),
    ("scenario", "Scenario", 32),
    ("preconditions", "Preconditions", 28),
    ("steps", "Steps", 40),
    ("test_data", "Test Data", 24),
    ("expected_result", "Expected Result", 32),
]


def _write_header(ws: Worksheet) -> None:
    for col_idx, (_, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"


def _write_test_cases_sheet(ws: Worksheet, test_cases: list[TestCase]) -> None:
    _write_header(ws)
    for row_idx, tc in enumerate(test_cases, start=2):
        values = {
            "tc_id": tc.tc_id,
            "module": tc.module,
            "case_type": tc.case_type.value,
            "priority": tc.priority.value,
            "scenario": tc.scenario,
            "preconditions": tc.preconditions,
            "steps": "\n".join(f"{i}. {s}" for i, s in enumerate(tc.steps, start=1)),
            "test_data": tc.test_data,
            "expected_result": tc.expected_result,
        }
        for col_idx, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "priority":
                cell.fill = PRIORITY_FILLS.get(tc.priority.value, PatternFill())


def _write_summary_sheet(ws: Worksheet, test_cases: list[TestCase]) -> None:
    ws.cell(row=1, column=1, value="Test Case Generation Summary").font = Font(bold=True, size=14)

    by_type = Counter(tc.case_type.value for tc in test_cases)
    by_priority = Counter(tc.priority.value for tc in test_cases)
    by_module = Counter(tc.module for tc in test_cases)

    row = 3
    ws.cell(row=row, column=1, value="Total test cases").font = Font(bold=True)
    ws.cell(row=row, column=2, value=len(test_cases))
    row += 2

    ws.cell(row=row, column=1, value="By type").font = Font(bold=True)
    row += 1
    for case_type, count in sorted(by_type.items()):
        ws.cell(row=row, column=1, value=case_type)
        ws.cell(row=row, column=2, value=count)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="By priority").font = Font(bold=True)
    row += 1
    for priority, count in sorted(by_priority.items()):
        ws.cell(row=row, column=1, value=priority)
        ws.cell(row=row, column=2, value=count)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="By module").font = Font(bold=True)
    row += 1
    for module, count in sorted(by_module.items()):
        ws.cell(row=row, column=1, value=module)
        ws.cell(row=row, column=2, value=count)
        row += 1

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14


def build_workbook(test_cases: list[TestCase]) -> Workbook:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Summary"
    _write_summary_sheet(summary_ws, test_cases)

    tc_ws = wb.create_sheet("Test Cases")
    _write_test_cases_sheet(tc_ws, test_cases)

    return wb


def export_to_file(test_cases: list[TestCase], output_path: str) -> str:
    wb = build_workbook(test_cases)
    wb.save(output_path)
    return output_path
